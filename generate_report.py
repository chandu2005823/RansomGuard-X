import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()

# Title
title = doc.add_heading('RansomGuard-X Project Status Report', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.add_paragraph('This document outlines the exact state of the RansomGuard-X (Phase 5B) project, organized by what has been completed and what remains to be executed.')

# Completed Work Section
h1 = doc.add_heading('COMPLETED WORK (Architecture & Engineering)', level=1)

def add_item(doc, title, body):
    p = doc.add_paragraph()
    runner = p.add_run(title)
    runner.bold = True
    p.add_run('\n' + body)

add_item(doc, '1. Telemetry & Feature Extraction Stabilized', 
         '- Race Conditions Fixed: Rewrote FeatureExtractor and background monitors to enforce strict, synchronous 10-second polling windows (auto_flush=False).\n'
         '- Leakage Eliminated: Enforced total monitor teardown/recreation per window in the Live Pipeline to guarantee that dynamic features never leak between snapshots.')

add_item(doc, '2. Live-Domain NLP Representation Solved',
         '- Shift from Cuckoo to Live: Proved that historical Cuckoo boolean flags cause severe hallucination (Domain Shift) on live Windows systems.\n'
         '- Phase 5B Converter: Engineered live_domain_converter.py to map the 21 raw telemetry integers into robust, categorical buckets (minimal, moderate, high-volume, extreme-volume).')

add_item(doc, '3. Data Collection Framework Built',
         '- Benign Engine: Created collect_benign_live.py and the interactive collect_benign.ps1 PowerShell menu to rapidly scale labeled background-noise datasets (e.g., benign_vscode, benign_browser).')

add_item(doc, '4. Machine Learning Training Pipeline Prepared',
         '- Trainer Script: Fully implemented train_live_bert.py using HuggingFace Trainer.\n'
         '- Leakage Prevention: Engineered a double GroupShuffleSplit on session_id to strictly guarantee no chronological data leaks across Train (70%), Val (15%), and Test (15%) sets.\n'
         '- Class Imbalance: Integrated dynamic class-weight calculation (compute_class_weight) mapped to a custom CrossEntropyLoss function.')

add_item(doc, '5. Inference & XAI Pipeline Prepared',
         '- Text Wrappers: Injected predict_text(text) into LiveBERTInference and explain_text(text) into BertXAI to natively parse Phase 5B generated text and yield LIME explanations.')

add_item(doc, '6. SOC/EDR Dashboard Redesign',
         '- Enterprise UI: Transformed app.py into a modern SOC dashboard prioritizing the Threat Workflow (AM I SAFE? -> MODEL STATUS -> OBSERVED BEHAVIOR).\n'
         '- UI Hardening: Fully severed the legacy models from the primary threat banner. The dashboard is safely locked to "ANALYSIS PENDING" until the new Live-Domain BERT is explicitly trained.')

# Pending Work Section
doc.add_page_break()
h2 = doc.add_heading('WORK YET TO DO (Data Execution & Training)', level=1)

add_item(doc, '1. Scale Benign Telemetry Collection',
         '- Action: Run collect_benign.ps1 across various normal Windows activities (e.g., coding, browsing, document editing, idle) to harvest hours of realistic benign baseline data.')

add_item(doc, '2. Execute Malicious Telemetry Collection (Phase 5B Detonation)',
         '- Action: Safely detonate ransomware (or simulated ransomware scripts) while running the identical Live Detection Pipeline to harvest the malicious [label = 1] telemetry mapping to the 21-feature schema.')

add_item(doc, '3. Compile & Validate the Dataset',
         '- Action: Merge the benign and malicious CSVs into a unified dataset.\n'
         '- Action: Run `python train_live_bert.py <merged_dataset.csv> validate` to ensure session_id isolation and class integrity are perfect.')

add_item(doc, '4. Train the Live-Domain BERT Model',
         '- Action: Execute `python train_live_bert.py <merged_dataset.csv> train`.\n'
         '- Outcome: This will fine-tune bert-base-uncased, evaluate against the untouch test set, and output the final model, tokenizer, and metrics to 6_experiments/live_domain_bert_model/.')

add_item(doc, '5. Connect the Dashboard',
         '- Action: Modify live_bert_inference.py to load the newly trained model directory instead of the historical one.\n'
         '- Outcome: The Streamlit dashboard\'s Threat Status will instantly unlock, shifting from ANALYSIS PENDING to live SYSTEM SAFE or RANSOMWARE DETECTED, populating LIME explanations dynamically.')

add_item(doc, '6. Final Live Verification',
         '- Action: Open the Dashboard and interact with the Windows environment live to confirm the model generalizes well to background noise and correctly flags suspicious behavior.')

doc.save('RansomGuard-X_Project_Status.docx')
