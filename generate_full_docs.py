import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()

# Title
title = doc.add_heading('RansomGuard-X: Comprehensive Project Documentation', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.add_paragraph('AI-Powered Ransomware Detection & Response System\nGenerated via MLRAN Phase 5')

# 1. Introduction
doc.add_heading('1. Project Overview & Objective', level=1)
doc.add_paragraph(
    'RansomGuard-X (also known internally as MLRAN) is an advanced endpoint detection and response (EDR) platform '
    'designed to detect ransomware behavior on live Windows environments. Instead of relying on traditional static signatures, '
    'the system tracks live system telemetry in 10-second sliding windows, translates raw mathematical features into '
    'human-readable behavioral text, and feeds this text into a fine-tuned NLP model (BERT) for classification.'
)

# 2. System Architecture
doc.add_heading('2. Core Architecture', level=1)

def add_bullet(doc, title, body):
    p = doc.add_paragraph(style='List Bullet')
    runner = p.add_run(title)
    runner.bold = True
    p.add_run(': ' + body)

add_bullet(doc, 'Live Telemetry Engine', 'Monitors Windows processes, file systems, registries, network connections, and DLL loads using psutil and custom OS-level polling. Extracts data synchronously every 10 seconds.')
add_bullet(doc, 'Live Domain Converter', 'A deterministic translation layer that maps the 21 numerical features into categorical text (e.g., "extreme-volume file modifications"). This eliminates the "domain shift" problem caused by raw integers destroying NLP token embeddings.')
add_bullet(doc, 'BERT Inference Engine', 'A HuggingFace Transformer model (bert-base-uncased) fine-tuned for Sequence Classification. It reads the behavioral text and outputs a Benign or Ransomware probability.')
add_bullet(doc, 'Explainable AI (LIME)', 'Uses the Local Interpretable Model-agnostic Explanations (LIME) framework to highlight the specific text phrases (behaviors) that drove the model\'s decision, providing SOC analysts with transparent attribution.')
add_bullet(doc, 'Response Mitigator', 'An automated defense script capable of instantly suspending or resuming malicious Process IDs (PIDs) upon detection. Currently restricted to Dry-Run mode for system safety.')

# 3. Data Schema (21 Features)
doc.add_heading('3. The 21-Feature Telemetry Schema', level=1)
doc.add_paragraph('The pipeline extracts 21 features across 8 major OS categories per 10-second snapshot:')
features = [
    "File System (FS): Creations, Modifications, Deletions, Renames, Unique Extensions, Executable Drops.",
    "Process (PR): Child Processes, Suspicious Children, Terminations.",
    "Registry (RG): Keys Created, Values Modified, Persistence Modifications.",
    "Network (NW): Outbound Connections, Unique Destination IPs.",
    "Memory (MEM): Remote Threads Created, Cross-Process Access.",
    "Module (DLL): Image Loads, Unsigned Image Loads.",
    "Relationship (REL): Vulnerable Parent Process Spawning.",
    "Temporal (TMP): File Operations per second, Registry Operations per second."
]
for f in features:
    doc.add_paragraph(f, style='List Bullet')

# 4. Codebase Overview
doc.add_heading('4. Codebase Structure', level=1)
files = {
    'app.py': 'The primary Streamlit SOC/EDR dashboard. Provides visual threat status, grouped telemetry, XAI insights, and mitigation controls.',
    'live_detection_pipeline.py': 'Orchestrates the 10-second sliding window. Initializes monitors, pauses, flushes data synchronously, and passes the vector to the converter.',
    'feature_extractor.py': 'Core aggregation engine. Maps OS events to the 21-feature schema for active Process IDs.',
    'live_domain_converter.py': 'NLP translation logic. Buckets numerical volumes (minimal, moderate, high, extreme) and maps them into Phase 5B behavioral text.',
    'live_bert_inference.py': 'Wrapper for the HuggingFace BERT model. Contains predict_text() to output logits and probabilities.',
    'bert_xai.py': 'Wrapper for LIME text explainer. Integrates directly with BERT inference to provide phrase-level attribution.',
    'train_live_bert.py': 'The Phase 5B HuggingFace Trainer script. Includes leakage-proof Session ID splitting, class imbalance handling, and evaluation metrics.',
    'collect_benign_live.py & collect_benign.ps1': 'Automated collection scripts to harvest hours of live benign Windows telemetry to train the baseline.'
}
for name, desc in files.items():
    add_bullet(doc, name, desc)

# 5. Project Phases & Evolution
doc.add_heading('5. Project Phases & Evolution', level=1)
doc.add_paragraph(
    'Phase 5A (Historical Cuckoo Dataset): The project initially utilized a legacy Random Forest (tabular) model and an '
    '11-feature BERT model trained on historical Cuckoo Sandbox boolean flags. It was discovered that applying this sandbox '
    'model to a live Windows environment resulted in a catastrophic Domain Shift, where normal benign background noise '
    '(e.g., a single background process or connection) was incorrectly flagged as ransomware.'
)
doc.add_paragraph(
    'Phase 5B (Live-Domain Migration): The current phase. The system was completely overhauled to rely on Live-Domain categorical '
    'bucketing. The pipeline now natively handles the 21-feature telemetry and parses it into NLP sentences. We are currently '
    'in the data collection phase, harvesting real benign Windows activity and safe malicious detonations to train the final '
    'Live-Domain BERT model.'
)

# 6. Usage & SOC Dashboard
doc.add_heading('6. Usage Instructions', level=1)
doc.add_paragraph('To start the SOC Dashboard UI:')
p = doc.add_paragraph('streamlit run app.py')
p.style = 'Intense Quote'
doc.add_paragraph('The dashboard provides two modes:')
add_bullet(doc, 'Replay Validation Data', 'Cycles through historical CSV datasets for UI testing and model validation.')
add_bullet(doc, 'Live Benign Telemetry', 'Captures a real 10-second snapshot of the host machine, processes the telemetry, generates behavioral text, and feeds it into the detection pipeline live.')

doc.add_paragraph('\n-- End of Documentation --')
doc.save('RansomGuard-X_Full_Documentation.docx')
